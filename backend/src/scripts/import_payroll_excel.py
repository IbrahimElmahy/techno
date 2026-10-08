from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from decimal import Decimal

from sqlalchemy import select

from src.core.db import SessionLocal
from src.core.money import to_money
from src.models.employee import Employee
from src.models.hr_payroll import ComponentKind, EmployeeSalaryLine, SalaryComponent
from src.models.hr_payroll_sheet import PayrollGroup, PayrollGroupMember
from src.models.org import Branch
from src.services import numbering
from src.services import payroll_setup_service as setup
from src.services import payroll_sheet_service as sheet

GROUP_NAMES = ["البيع", "الإدارية", "خدمة العملاء"]

FIXED = {
    "اساسى": "__basic__",
    "بدلات": ("بدلات", "بدلات اضافية"),
    "تقييم": "تقييم",
    "منحة غلاء": "منحة غلاء",
    "مكالمات": "مكالمات",
    "اشراف": "اشراف",
    "تأمينات": "تأمينات",
}
TOTALS = {"الاجمالي", "الاجمالى", "الإجمالي", "الصافي", "الصافى"}


def norm(text) -> str:
    s = str(text or "")
    s = re.sub(r"[\u064B-\u065F\u0670]", "", s)
    s = s.replace("ـ", "")
    s = re.sub(r"[أإآٱ]", "ا", s)
    s = s.replace("ى", "ي").replace("ة", "ه").replace("ؤ", "و").replace("ئ", "ي")
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r"\bعبد\s+", "عبد", s)
    return s


def _num(v) -> Decimal | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return to_money(Decimal(str(v)))
    try:
        return to_money(Decimal(str(v).strip()))
    except Exception:  # noqa: BLE001
        return None


def read_tables(path: str, sheet_name: str | None):
    import openpyxl

    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    names = [sheet_name] if sheet_name else values.sheetnames
    for name in names:
        ws, wf = values[name], formulas[name]
        tables = []
        r = 1
        while r <= ws.max_row:
            labels = {c: str(ws.cell(r, c).value).strip() for c in range(3, ws.max_column + 1)
                      if ws.cell(r, c).value is not None}
            if ws.cell(r, 3).value is not None and str(ws.cell(r, 3).value).strip() == "اساسى":
                header = {c: v for c, v in labels.items()}
                rows = []
                r += 1
                while r <= ws.max_row:
                    who = ws.cell(r, 2).value
                    if who is None or str(who).strip() in TOTALS:
                        break
                    cells = {}
                    for c in header:
                        raw = wf.cell(r, c).value
                        cells[c] = (ws.cell(r, c).value,
                                    isinstance(raw, str) and raw.startswith("="))
                    rows.append((str(who).strip(), cells))
                    r += 1
                tables.append((header, rows))
            r += 1
        if tables:
            return name, tables
    raise SystemExit("مالقيتش جدول مرتبات (صف عناوين أوله «اساسى») في الملف.")


def match(db, branch_id: int, names: list[str], manual: dict[str, str]):
    emps = db.scalars(select(Employee).where(Employee.branch_id == branch_id,
                                             Employee.active.is_(True))).all()
    others = db.scalars(select(Employee).where(Employee.branch_id != branch_id)).all()
    by_norm: dict[str, list[Employee]] = {}
    for e in emps:
        by_norm.setdefault(norm(e.name), []).append(e)
    out = {}
    for name in names:
        n = norm(name)
        if name in manual:
            e = db.scalar(select(Employee).where(Employee.code == manual[name]))
            out[name] = ("manual", e) if e else ("unmatched", None)
            continue
        exact = by_norm.get(n, [])
        if len(exact) == 1:
            out[name] = ("exact", exact[0])
            continue
        if len(exact) > 1:
            out[name] = ("ambiguous", exact)
            continue
        words = n.split()
        probable = [e for e in emps if norm(e.name).split()[:len(words)] == words]
        if len(probable) == 1:
            out[name] = ("probable", probable[0])
            continue
        elsewhere = [e for e in others if norm(e.name) == n]
        out[name] = ("unmatched", elsewhere)
    return out


def run(args) -> None:
    db = SessionLocal()
    try:
        sheet_name, tables = read_tables(args.file, args.sheet)
        branches = db.scalars(select(Branch)).all()
        if args.branch:
            branch = next((b for b in branches if norm(args.branch) in norm(b.name)), None)
            if branch is None:
                raise SystemExit(f"الفرع «{args.branch}» مش موجود.")
        else:
            names = {norm(n) for _, rows in tables[:3] for n, _ in rows}
            score = {}
            for b in branches:
                emps = db.scalars(select(Employee.name).where(
                    Employee.branch_id == b.id, Employee.active.is_(True))).all()
                score[b.id] = len(names & {norm(e) for e in emps})
            branch = max(branches, key=lambda b: score[b.id])
            print("الفرع اتعرف من الأسماء: " + "، ".join(
                f"{b.name}={score[b.id]}" for b in branches))
        print(f"الملف: {args.file}\nالورقة: {sheet_name} — {len(tables)} جدول\n"
              f"الفرع: {branch.name} (#{branch.id})\n")
        if len(tables) > len(GROUP_NAMES):
            print(f"⚠️ فيه {len(tables)} جدول — أول {len(GROUP_NAMES)} بس هيتقروا.")

        actor = args.actor_user_id
        made = sheet.apply_template(db, branch_id=branch.id, actor_user_id=actor)
        print("المجموعات:", "اتعملت " + "، ".join(g.name for g in made) if made
              else "موجودة بالفعل")
        groups = {g.name: g for g in db.scalars(select(PayrollGroup).where(
            PayrollGroup.branch_id == branch.id)).all()}
        comps = {c.name: c for c in db.scalars(select(SalaryComponent)).all()}

        manual = {}
        for item in args.map or []:
            if "=" in item:
                k, v = item.split("=", 1)
                manual[k.strip()] = v.strip()

        all_names = [n for _, rows in tables[:3] for n, _ in rows]
        matches = match(db, branch.id, all_names, manual)
        eff = date.fromisoformat(args.effective_from) if args.effective_from else \
            date.today().replace(day=1)

        stats = {"assigned": 0, "salary": 0, "created": 0, "skipped": 0}
        monthly: list[str] = []
        for gi, (header, rows) in enumerate(tables[:3]):
            gname = GROUP_NAMES[gi]
            group = groups[gname]
            print(f"\n=== {gname} ({len(rows)} صف) ===")
            seen_label: dict[str, int] = {}
            col_plan = {}
            for c, label in header.items():
                n = seen_label.get(label, 0)
                seen_label[label] = n + 1
                target = FIXED.get(label)
                if isinstance(target, tuple):
                    target = target[min(n, len(target) - 1)]
                col_plan[c] = (label, target)
            order = 0
            for name, cells in rows:
                status, who = matches[name]
                emp = None
                if status in ("exact", "manual"):
                    emp = who
                elif status == "probable" and args.accept_probable:
                    emp = who
                elif status in ("unmatched", "probable") and args.create_missing:
                    emp = Employee(
                        code=numbering.next_document_number(
                            db, Employee, "EMP", column=Employee.code, width=4),
                        name=name, branch_id=branch.id,
                        notes="اتعمل من استيراد ورقة المرتبات")
                    db.add(emp)
                    db.flush()
                    stats["created"] += 1
                    status = "created"
                tag = {"exact": "✓", "manual": "✓ يدوي", "probable": "؟ محتمل",
                       "created": "+ جديد", "ambiguous": "✗ أكتر من موظف",
                       "unmatched": "✗ مش موجود"}[status]
                extra = ""
                if status == "probable" and not args.accept_probable:
                    extra = f" ← {who.name} ({who.code}) — مش متطبّق (--accept-probable)"
                elif status == "ambiguous":
                    extra = " ← " + "، ".join(f"{e.name} ({e.code})" for e in who)
                elif status == "unmatched" and who:
                    extra = " — موجود في فرع تاني: " + "، ".join(
                        f"{e.name} ({e.code}، فرع {e.branch_id})" for e in who)
                elif emp is not None and status != "created":
                    extra = f" ← {emp.name} ({emp.code})"
                if emp is None:
                    print(f"  {tag}  {name}{extra}")
                    stats["skipped"] += 1
                    continue

                order += 1
                sheet.assign(db, branch_id=branch.id, employee_ids=[emp.id],
                             group_id=group.id, actor_user_id=actor)
                m = db.scalar(select(PayrollGroupMember).where(
                    PayrollGroupMember.employee_id == emp.id))
                m.sort_order = order
                stats["assigned"] += 1

                basic = None
                lines: dict[int, Decimal] = {}
                month_vals = []
                for c, (label, target) in col_plan.items():
                    value, is_formula = cells.get(c, (None, False))
                    num = _num(value)
                    if target is None:
                        if num and not is_formula and label not in TOTALS and label != "غياب":
                            month_vals.append(f"{label}={num}")
                        continue
                    if num is None or is_formula:
                        continue
                    if target == "__basic__":
                        basic = num
                    elif num:
                        comp = comps.get(target) or sheet.ensure_component(
                            db, name=target, kind=ComponentKind.deduction if target == "تأمينات"
                            else ComponentKind.earning, actor_user_id=actor)
                        comps[target] = comp
                        lines[comp.id] = num
                if month_vals:
                    monthly.append(f"{emp.name}: " + "، ".join(month_vals))
                if basic is None and not lines:
                    print(f"  {tag}  {name}{extra} — مافيش أرقام ثابتة في الملف")
                    continue
                current = setup.salary_on(db, emp.id, eff)
                merged: dict[int, dict] = {}
                if current is not None:
                    for ln in db.scalars(select(EmployeeSalaryLine).where(
                            EmployeeSalaryLine.salary_id == current.id)).all():
                        merged[ln.component_id] = {"component_id": ln.component_id,
                                                   "amount": ln.amount, "pct": ln.pct}
                for cid, amount in lines.items():
                    merged[cid] = {"component_id": cid, "amount": amount, "pct": None}
                final_basic = basic if basic is not None else (
                    current.basic if current is not None else 0)
                setup.set_salary(
                    db, employee_id=emp.id, effective_from=eff, basic=final_basic,
                    actor_user_id=actor,
                    insurance_base=current.insurance_base if current is not None else None,
                    lines=list(merged.values()),
                    notes="من ورقة المرتبات (إكسل)")
                stats["salary"] += 1
                desc = "، ".join(f"{n}={v}" for n, v in (
                    [("اساسى", final_basic)] + [(db.get(SalaryComponent, k).name, v)
                                                for k, v in lines.items()]))
                print(f"  {tag}  {name}{extra} — {desc}")

        if monthly:
            print("\nأرقام شهرية في الملف (مش إعدادات — بتتكتب في شيت الشهر):")
            for line in monthly:
                print("  •", line)
        print(f"\nالخلاصة: اتوزّع {stats['assigned']}، إعدادات راتب {stats['salary']}، "
              f"موظفين جداد {stats['created']}، متساب {stats['skipped']}")
        if args.yes:
            db.commit()
            print("✓ اتحفظ.")
        else:
            db.rollback()
            print("(تجربة — مااتحفظش حاجة. ضيف --yes للتنفيذ)")
    finally:
        db.close()


def main(argv=None) -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--file", required=True)
    p.add_argument("--sheet")
    p.add_argument("--branch", help="افتراضي: الفرع اللي فيه أكتر أسماء متطابقة")
    p.add_argument("--effective-from")
    p.add_argument("--map", action="append")
    p.add_argument("--accept-probable", action="store_true")
    p.add_argument("--create-missing", action="store_true")
    p.add_argument("--actor-user-id", type=int, default=None)
    p.add_argument("--yes", action="store_true")
    args = p.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    run(args)


if __name__ == "__main__":
    main()
