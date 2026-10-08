from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from decimal import Decimal

from sqlalchemy import select

from src.core.db import SessionLocal
from src.models.employee import Employee
from src.models.hr_commission import (
    HrCommissionSupervisor,
    HrCommissionTeam,
    HrCommissionTechnician,
)
from src.models.org import Branch
from src.services import hr_commission_service as svc

DEFAULT_XLSX = (r"C:\Users\Ibrahim Elmahy\.claude\uploads\2cb30e42-91df-46f3-899d-a0f62e35aa2e"
                r"\b66a6b0b-Book2.xlsx_____________.xlsx")
DEFAULT_SHEET = "اكتوبر (2)"
DEFAULT_BRANCH = "العلياء"

TEAMS = (
    ("سيارة أ", "G94", "J94", (8, 17), True, "C66", True),
    ("سيارة ب", "G131", "J131", (11, 12), True, "C55", True),
    ("سيارة ج", "G112", "J112", (7, 6), True, "G66", True),
    ("سيارة د", "G150", "J150", (15, 16), True, "G55", True),
    ("الشرقية", "F80", "I80", (14, 13), False, None, False),
    ("اسكندرية", None, None, (), True, None, False),
    ("الضبعة", None, None, (), True, None, False),
)
PENALIZED_ROWS = {8, 11, 7, 6, 15}

SUPERVISORS = (
    (4, "اشراف عام", "Q11", "Q11",
     ("سيارة أ", "سيارة ب", "سيارة ج", "سيارة د", "الشرقية", "اسكندرية", "الضبعة"),
     1, False, True, None),
    (8, "اشراف على أ/ج", "O45", "O45", ("سيارة أ", "سيارة ج"), 1, True, False, None),
    (9, "اشراف على ب/د", "P54", "O54", ("سيارة ب", "سيارة د"), 0, True, True, "H41"),
    (10, "١٪ من اسكندرية", "E10", "E10", ("اسكندرية",), 1, False, True, None),
)

TECH_ROWS = range(167, 174)
TECH_BONUS_ROWS = range(191, 197)
TECH_ALL_ROW = 173

_NUM = re.compile(r"(\d+(?:\.\d+)?)(%?)")


def _load(path: str, sheet: str):
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("محتاج openpyxl: pip install openpyxl") from exc
    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wf.sheetnames:
        raise SystemExit(f"الصفحة «{sheet}» مش في الملف. الموجود: {wf.sheetnames}")
    return wf[sheet], wv[sheet]


def _pct(wf, wv, ref: str | None) -> Decimal:
    if ref is None:
        return Decimal("0")
    raw = wf[ref].value
    if isinstance(raw, (int, float)):
        return (Decimal(str(raw)) * 100).normalize()
    text = str(raw or "")
    if text.startswith("="):
        m = re.search(r"(\d+(?:\.\d+)?)\s*/\s*100", text)
        if m:
            return Decimal(m.group(1))
        nums = _NUM.findall(text)
        pcts = [n for n, p in nums if p]
        if pcts:
            return Decimal(pcts[-1])
    cached = wv[ref].value
    if isinstance(cached, (int, float)):
        return (Decimal(str(cached)) * 100).normalize()
    raise SystemExit(f"مش عارف أقرا نسبة من {ref}: {raw!r}")


def _num(wf, wv, ref: str) -> Decimal:
    v = wv[ref].value
    if v is None:
        v = wf[ref].value
    try:
        return Decimal(str(v)) if v not in (None, "") else Decimal("0")
    except Exception:  # noqa: BLE001
        return Decimal("0")


def _const_after(wf, ref: str, op: str) -> Decimal | None:
    text = str(wf[ref].value or "")
    m = re.search(re.escape(op) + r"\s*(\d+(?:\.\d+)?)", text)
    return Decimal(m.group(1)) if m else None


_FOLD = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ى": "ي", "ة": "ه", "ـ": "",
                       "\xa0": " ", ".": " ", "،": " ", "-": " ", "(": " ", ")": " "})


def name_key(name: str) -> str:
    s = unicodedata.normalize("NFKC", str(name or "")).translate(_FOLD)
    s = "".join(c for c in s if not ("\u064b" <= c <= "\u0652"))
    s = re.sub(r"ي{2,}", "ي", s)
    return " ".join(s.split()).casefold()


def match_employee(name: str, employees: list[Employee]) -> tuple[Employee | None, str]:
    key = name_key(name)
    if not key:
        return None, "اسم فاضي"
    tokens = key.split()
    tiers = (
        ("مطابق", lambda k: k == key),
        ("مطابق من غير مسافات", lambda k: k.replace(" ", "") == key.replace(" ", "")),
        ("الاسم أوله", lambda k: k.startswith(key + " ")),
        ("كلماته جوّاه", lambda k: all(t in k.split() for t in tokens)),
    )
    for label, test in tiers:
        hits = [e for e in employees if test(name_key(e.name))]
        if not hits:
            continue
        active = [e for e in hits if e.active]
        pool = active or hits
        if len(pool) == 1:
            return pool[0], label
        return None, "ملتبس: " + " / ".join(f"{e.name} [{e.id}]" for e in pool)
    return None, "مالوش موظف"


def build_plan(wf, wv, employees: list[Employee], picks: dict[str, int] | None = None) -> dict:
    report: list[str] = []
    unmatched: list[str] = []
    by_id = {e.id: e for e in employees}
    picks = {name_key(k): v for k, v in (picks or {}).items()}

    def emp(name: str, where: str) -> Employee | None:
        if name_key(name) in picks:
            e = by_id.get(picks[name_key(name)])
            if e is None:
                unmatched.append(f"{where}: «{name}» — --pick بيشاور على موظف مش في الفرع")
            else:
                report.append(f"{where}: «{name}» → «{e.name}» [{e.id}] (--pick)")
            return e
        e, how = match_employee(name, employees)
        if e is None:
            unmatched.append(f"{where}: «{name}» — {how}")
        elif how != "مطابق":
            report.append(f"{where}: «{name}» → «{e.name}» [{e.id}] ({how})")
        return e

    def row_name(r: int) -> str:
        return " ".join(str(wf[f"B{r}"].value or "").split())

    settings = {
        "point_value": _const_after(wf, "E167", "*") or Decimal("0"),
        "points_per_coupon": int(_const_after(wf, "D167", "*") or 30),
        "factor_divisor": _const_after(wf, "G167", "/") or Decimal("600"),
        "absence_divisor": _const_after(wf, "H99", "/") or Decimal("30"),
        "penalty_sales_pct": _pct(wf, wv, "C53") if wf["C53"].value else Decimal("25"),
        "penalty_per_thousand": _const_after(wf, "C58", "*") or Decimal("20"),
        "attribute_by_customer_rep": True,
        "pay_coupon_commission": False,
        "apply_min_inspections": False,
    }

    teams = []
    for name, poly_ref, white_ref, rows, split, credit_ref, penalty in TEAMS:
        rate_poly = _pct(wf, wv, poly_ref)
        rate_white = _pct(wf, wv, white_ref)
        members = []
        for r in rows:
            e = emp(row_name(r), f"{name} (صف {r})")
            if e is not None:
                members.append({"employee": e, "exempt_25": r not in PENALIZED_ROWS})
        teams.append({
            "name": name, "rate_poly": rate_poly, "rate_white": rate_white,
            "rate_other": rate_white, "split_equally": split, "penalty_enabled": penalty,
            "credit_limit": _num(wf, wv, credit_ref) if credit_ref else Decimal("0"),
            "members": members,
            "users": sorted({m["employee"].user_id for m in members if m["employee"].user_id}),
            "notes": ("تحصيل يدوي — الرقم في الملف مكتوب بالإيد" if not rows else None),
        })

    sups = []
    for r, label, poly_ref, white_ref, tnames, offset, absence, active, pen_ref in SUPERVISORS:
        e = emp(row_name(r), f"مشرف «{label}» (صف {r})")
        if e is None:
            continue
        rp, rw = _pct(wf, wv, poly_ref), _pct(wf, wv, white_ref)
        sups.append({
            "employee": e, "label": label, "rate_poly": rp, "rate_white": rw, "rate_other": rw,
            "teams": tnames, "period_offset": offset, "deduct_absence": absence,
            "active": active, "penalty_rate": _pct(wf, wv, pen_ref) if pen_ref else Decimal("0"),
            "notes": (None if active else "محسوب في الملف (O48) ومش داخل المرتب — موقوف"),
        })

    bonus_rate = {}
    for r in TECH_BONUS_ROWS:
        n = name_key(wf[f"E{r}"].value or "")
        if n:
            bonus_rate[n] = _num(wf, wv, f"J{r}")
    techs = []
    for r in TECH_ROWS:
        raw = " ".join(str(wf[f"B{r}"].value or "").split())
        if not raw:
            continue
        e = emp(raw, f"فني (صف {r})")
        if e is None:
            continue
        techs.append({
            "employee": e, "excel_name": raw, "factor": _num(wf, wv, f"F{r}"),
            "min_inspections": int(_num(wf, wv, f"K{r}")),
            "inspection_rate": _num(wf, wv, f"N{r}"),
            "plumber_rate": bonus_rate.get(name_key(raw), Decimal("0")),
            "scope": "all" if r == TECH_ALL_ROW else "own",
            "sort_order": r - TECH_ROWS.start,
        })
    return {"settings": settings, "teams": teams, "supervisors": sups, "technicians": techs,
            "report": report, "unmatched": unmatched}


def _print_plan(plan: dict, branch: Branch) -> None:
    p = print
    p(f"\n=== إعدادات العمولات لفرع «{branch.name}» [{branch.id}] ===")
    p("\nأرقام الفرع:")
    for k, v in plan["settings"].items():
        p(f"  {k}: {v}")
    p("\nالسيارات:")
    for t in plan["teams"]:
        p(f"  • {t['name']}: بولي {t['rate_poly']}٪ / أبيض {t['rate_white']}٪ / من غير عيلة "
          f"{t['rate_other']}٪ — {'بالتساوي' if t['split_equally'] else 'كاملة لكل فرد'}"
          + (f" — خصم ٢٥٪، ائتمان {t['credit_limit']}" if t["penalty_enabled"] else ""))
        for m in t["members"]:
            e = m["employee"]
            p(f"      - {e.name} [{e.id}]"
              + (" (معفي من ٢٥٪)" if m["exempt_25"] and t["penalty_enabled"] else ""))
        p(f"      حسابات المناديب: {t['users'] or '—'}")
    p("\nالمشرفين:")
    for s in plan["supervisors"]:
        p(f"  • {s['employee'].name} [{s['employee'].id}] «{s['label']}»: بولي {s['rate_poly']}٪"
          f" / أبيض {s['rate_white']}٪ على {', '.join(s['teams'])}"
          f" — {'الشهر اللي فات' if s['period_offset'] else 'نفس الشهر'}"
          + (", بخصم غياب" if s["deduct_absence"] else "")
          + (f", خصم {s['penalty_rate']}٪ من زيادة الـ٢٥٪" if s["penalty_rate"] else "")
          + ("" if s["active"] else "  [موقوف]"))
    p("\nالفنيين:")
    for t in plan["technicians"]:
        p(f"  • {t['employee'].name} [{t['employee'].id}] (في الملف «{t['excel_name']}»): "
          f"معامل {t['factor']}، حد أدنى {t['min_inspections']}، معاينة {t['inspection_rate']}، "
          f"سباك {t['plumber_rate']}" + ("، على كل الفنيين" if t["scope"] == "all" else "")
          + ("" if t["employee"].user_id else "  ⚠ مالوش حساب تطبيق"))
    if plan["report"]:
        p("\nمطابقات مش حرفية (راجعها):")
        for line in plan["report"]:
            p("  " + line)
    if plan["unmatched"]:
        p("\n⚠ أسماء مالهاش موظف (مش هتتسجّل — ضيفهم من الشاشة):")
        for line in plan["unmatched"]:
            p("  " + line)


def apply_plan(db, plan: dict, branch_id: int, actor_user_id: int | None) -> None:
    svc.save_settings(db, branch_id=branch_id, data=plan["settings"],
                      actor_user_id=actor_user_id)
    team_ids: dict[str, int] = {}
    for t in plan["teams"]:
        existing = db.scalar(select(HrCommissionTeam).where(
            HrCommissionTeam.branch_id == branch_id, HrCommissionTeam.name == t["name"]))
        data = {k: t[k] for k in ("name", "rate_poly", "rate_white", "rate_other",
                                  "split_equally", "penalty_enabled", "credit_limit", "notes")}
        data |= {"active": True, "sort_order": len(team_ids),
                 "users": t["users"],
                 "members": [{"employee_id": m["employee"].id, "exempt_25": m["exempt_25"]}
                             for m in t["members"]]}
        row = svc.save_team(db, branch_id=branch_id, data=data, actor_user_id=actor_user_id,
                            team_id=existing.id if existing else None)
        team_ids[t["name"]] = row.id
    for i, s in enumerate(plan["supervisors"]):
        existing = db.scalar(select(HrCommissionSupervisor).where(
            HrCommissionSupervisor.branch_id == branch_id,
            HrCommissionSupervisor.employee_id == s["employee"].id,
            HrCommissionSupervisor.label == s["label"]))
        data = {k: s[k] for k in ("label", "rate_poly", "rate_white", "rate_other",
                                  "penalty_rate", "period_offset", "deduct_absence", "active",
                                  "notes")}
        data |= {"employee_id": s["employee"].id, "sort_order": i,
                 "teams": [team_ids[n] for n in s["teams"] if n in team_ids]}
        svc.save_supervisor(db, branch_id=branch_id, data=data, actor_user_id=actor_user_id,
                            supervisor_id=existing.id if existing else None)
    for t in plan["technicians"]:
        existing = db.scalar(select(HrCommissionTechnician).where(
            HrCommissionTechnician.branch_id == branch_id,
            HrCommissionTechnician.employee_id == t["employee"].id))
        data = {k: t[k] for k in ("factor", "min_inspections", "inspection_rate", "plumber_rate",
                                  "scope", "sort_order")}
        data |= {"employee_id": t["employee"].id, "active": True}
        svc.save_technician(db, branch_id=branch_id, data=data, actor_user_id=actor_user_id,
                            technician_id=existing.id if existing else None)


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--sheet", default=DEFAULT_SHEET)
    ap.add_argument("--branch", default=DEFAULT_BRANCH, help="اسم الفرع")
    ap.add_argument("--pick", action="append", default=[], metavar="الاسم=رقم_الموظف",
                    help="يحسم اسم ملتبس: --pick \"حسن رمضان=79\"")
    ap.add_argument("--yes", action="store_true", help="نفّذ (من غيرها عرض بس)")
    args = ap.parse_args(argv)
    picks = {}
    for item in args.pick:
        name, _, num = item.rpartition("=")
        if not name or not num.strip().isdigit():
            raise SystemExit(f"--pick لازم يبقى «الاسم=رقم»: {item!r}")
        picks[name.strip()] = int(num)

    wf, wv = _load(args.xlsx, args.sheet)
    db = SessionLocal()
    try:
        from src.core.db import Base, engine

        Base.metadata.create_all(engine)
        branch = db.scalar(select(Branch).where(Branch.name == args.branch))
        if branch is None:
            raise SystemExit(f"مافيش فرع اسمه «{args.branch}».")
        employees = db.scalars(select(Employee).where(Employee.branch_id == branch.id)).all()
        plan = build_plan(wf, wv, list(employees), picks)
        _print_plan(plan, branch)
        if not args.yes:
            print("\n(عرض بس — ضيف --yes للتنفيذ)")
            return
        apply_plan(db, plan, branch.id, actor_user_id=None)
        db.commit()
        print("\n✔ اتسجّل.")
    finally:
        db.close()


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
