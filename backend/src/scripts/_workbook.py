"""قراءة صفحات ملف العميل (الإكسل) مباشرةً — بدل خطوة التصدير اليدوي لـTSV.

سكربتات التكميل كانت بتقرا `owners.tsv` و`plumbers.tsv` من فولدر، والملفات دي
بتتعمل بإيد من الإكسل: تفتح الصفحة، تصدّرها، تسمّي الأعمدة بأسماء إنجليزي صغيرة.
الخطوة دي هي اللي بتتنسى — فالسكربت يتساب من غير ما يتشغّل، والأرقام اللي في
الملف تفضل في الملف. `--xlsx` بيشيلها: الملف هو المصدر زي ما العميل بعته.

**الأسماء بتتترجم هنا مرة واحدة.** رؤوس الملف بالعربي وبأسماء a5 (`Phone1`،
«فون»، `Address2`)، والسكربتات شايلة أسماء موحّدة (`phone1`، `floor`). الجدول
تحت هو الترجمة — مكتوبة في مكان واحد عشان تصليحها يبقى في مكان واحد.

⛔ الملف ده **هو** مصدر ما بعد البيع الوحيد (شوف `CLAUDE.md`) — قاعدة `ERP` ممنوعة.
"""
from __future__ import annotations

import os

#: صفحة «عملاء» = الملّاك. المفتاح اسم العمود في الملف، والقيمة الاسم اللي
#: السكربت بيدوّر عليه.
OWNERS_SHEET = "عملاء"
OWNERS_COLUMNS = {
    "ID": "id",
    "Code": "code",
    "Name_Ar": "name",
    "Phone1": "phone1",
    "Phone2": "phone2",
    "Mobile": "mobile",
    "Address": "address",
    # الاسم بيكدب والمحتوى هو اللي بيحدد: «ارضي»، «أول علوي» — ده الدور.
    "Address2": "floor",
}

PLUMBERS_SHEET = "سباك"
PLUMBERS_COLUMNS = {
    "Code": "code",
    "Name_Ar": "name",
    "النوع": "kind",
    "فون": "phone",
    "فون2": "phone2",
    "فون3": "phone3",
    "المحافظه": "gov",
    "المندينة": "city",
    "المنطقه": "area",
    "رقم المندوب": "rep_code",
}

_EMPTY = {"", "none", "null", "nan"}


def _clean(v: object) -> str:
    """القيمة كنص مقصوص — و«NULL» و«None» بيرجعوا فاضي.

    الملف مصدّر من SQL Server، فالفاضي بيوصل كلمة `NULL` مكتوبة. من غير التنضيف
    ده كنا هنكتب «NULL» في خانة التليفون ونفتكرها رقم.
    """
    s = "" if v is None else str(v).strip()
    return "" if s.lower() in _EMPTY else s


def read_sheet(path: str, sheet: str, columns: dict[str, str]) -> list[dict[str, str]]:
    """صفوف صفحة واحدة بالأسماء الموحّدة. العمود اللي مش في الملف بيرجع فاضي.

    الصف اللي كل خاناته فاضية بيتشال — الإكسل بيسيب صفوف فاضية في آخر الصفحة،
    وعدّها بيخلّي التقرير يقول أرقام مش حقيقية.
    """
    if not os.path.exists(path):
        raise SystemExit(f"مافيش الملف: {path}")
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - بيئة ناقصة
        raise SystemExit("محتاج openpyxl: pip install openpyxl") from exc

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"مافيش صفحة «{sheet}» في الملف — الصفحات: {wb.sheetnames}")
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    header = [_clean(h) for h in next(rows, ())]
    index = {out: header.index(src) for src, out in columns.items() if src in header}

    out: list[dict[str, str]] = []
    for r in rows:
        row = {name: _clean(r[i]) if i < len(r) else "" for name, i in index.items()}
        for name in columns.values():
            row.setdefault(name, "")
        if any(row.values()):
            out.append(row)
    wb.close()
    return out


def owners(path: str) -> list[dict[str, str]]:
    return read_sheet(path, OWNERS_SHEET, OWNERS_COLUMNS)


def plumbers(path: str) -> list[dict[str, str]]:
    return read_sheet(path, PLUMBERS_SHEET, PLUMBERS_COLUMNS)
