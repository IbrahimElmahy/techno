from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, Query, status
from pydantic import BaseModel
from sqlalchemy.orm import Session

from src.auth import branch_scope
from src.auth.dependencies import CurrentUser, require_capability
from src.auth.rbac import CAP_PAYROLL_POST, CAP_SALARY_VIEW
from src.core.db import get_db
from src.models.hr_payroll_run import PayrollRun
from src.models.hr_payroll_sheet import PayrollGroup
from src.services import payroll_sheet_service as sheet
from src.services.ledger_service import LedgerError
from src.services.payroll_service import PayrollError
from src.services.payroll_sheet_service import PayrollSheetError

router = APIRouter(tags=["payroll-sheet"], prefix="/hr/payroll-sheet")

_ERRORS = (PayrollSheetError, PayrollError, LedgerError)


_DIAC = re.compile("[\u064b-\u0652\u0670]")


def _has(text: str, needle: str) -> bool:
    return _DIAC.sub("", needle) in _DIAC.sub("", text)


def _raise(exc: Exception):
    text = str(exc)
    if isinstance(exc, LedgerError):
        raise HTTPException(409, {"code": "ledger_invalid", "message": text}) from exc
    if _has(text, "غير موجود"):
        raise HTTPException(404, {"code": "not_found", "message": text}) from exc
    if any(_has(text, s) for s in ("مرحّل", "معتمد", "متعكس", "معكوس")):
        raise HTTPException(409, {"code": "locked", "message": text}) from exc
    raise HTTPException(422, {"code": "validation", "message": text}) from exc


def _branch(current: CurrentUser, requested: int | None) -> int:
    if branch_scope.sees_all_branches(current):
        branch = requested if requested is not None else branch_scope.visible_branch_id(current)
        if branch is None:
            raise HTTPException(422, {"code": "validation", "message": "اختر الفرع."})
        return branch
    if requested is not None and requested != current.branch_id:
        raise HTTPException(404, {"code": "not_found", "message": "الفرع غير موجود."})
    if current.branch_id is None:
        raise HTTPException(422, {"code": "validation", "message": "اختر الفرع."})
    return current.branch_id


def _seen_run(db: Session, run_id: int, current: CurrentUser) -> PayrollRun:
    run = db.get(PayrollRun, run_id)
    ok = run is not None and sheet.is_sheet(db, run_id) and (
        branch_scope.sees_all_branches(current) or run.branch_id == current.branch_id)
    if not ok:
        raise HTTPException(404, {"code": "not_found", "message": "الشيت غير موجود."})
    return run


def _seen_group(db: Session, group_id: int, current: CurrentUser) -> PayrollGroup:
    g = db.get(PayrollGroup, group_id)
    if g is None or not (branch_scope.sees_all_branches(current)
                         or g.branch_id == current.branch_id):
        raise HTTPException(404, {"code": "not_found", "message": "المجموعة غير موجودة."})
    return g


class GroupIn(BaseModel):
    branch_id: int | None = None
    name: str
    columns: list[dict]
    absence_base: list[str] | None = None
    absence_divisor: int = 30
    sort_order: int | None = None
    active: bool = True
    notes: str | None = None


class BranchIn(BaseModel):
    branch_id: int | None = None


class CopyIn(BaseModel):
    from_branch_id: int
    to_branch_id: int | None = None


class OrderIn(BaseModel):
    branch_id: int | None = None
    ids: list[int]


class AssignIn(BaseModel):
    branch_id: int | None = None
    employee_ids: list[int]
    group_id: int | None = None


class MonthIn(BaseModel):
    branch_id: int | None = None
    year: int
    month: int


class CellIn(BaseModel):
    employee_id: int
    col_key: str
    value: Decimal | None = None


class DaysIn(BaseModel):
    employee_id: int
    days: Decimal | None = None


class NoteIn(BaseModel):
    employee_id: int
    notes: str | None = None


@router.get("/catalog")
def get_catalog(
    _: CurrentUser = Depends(require_capability(CAP_SALARY_VIEW)),
    db: Session = Depends(get_db),
) -> dict:
    return sheet.catalog(db)


@router.get("/groups")
def list_groups(
    branch_id: int | None = Query(None),
    include_inactive: bool = Query(False),
    current: CurrentUser = Depends(require_capability(CAP_SALARY_VIEW)),
    db: Session = Depends(get_db),
) -> list[dict]:
    return sheet.list_groups(db, branch_id=_branch(current, branch_id),
                             include_inactive=include_inactive)


@router.post("/groups", status_code=status.HTTP_201_CREATED)
def create_group(
    body: GroupIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    data = body.model_dump()
    data["branch_id"] = _branch(current, body.branch_id)
    try:
        g = sheet.save_group(db, actor_user_id=current.id, **data)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return {"id": g.id}


@router.put("/groups/{group_id}")
def update_group(
    group_id: int,
    body: GroupIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    g = _seen_group(db, group_id, current)
    data = body.model_dump()
    data["branch_id"] = g.branch_id
    try:
        sheet.save_group(db, actor_user_id=current.id, group_id=g.id, **data)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return {"id": g.id}


@router.delete("/groups/{group_id}")
def delete_group(
    group_id: int,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    _seen_group(db, group_id, current)
    try:
        name = sheet.delete_group(db, group_id=group_id, actor_user_id=current.id)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return {"deleted": name}


@router.post("/groups/order")
def order_groups(
    body: OrderIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    sheet.reorder_groups(db, branch_id=_branch(current, body.branch_id), group_ids=body.ids)
    db.commit()
    return {"ok": True}


@router.post("/groups/template")
def template(
    body: BranchIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    try:
        made = sheet.apply_template(db, branch_id=_branch(current, body.branch_id),
                                    actor_user_id=current.id)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return {"created": [g.name for g in made]}


@router.post("/groups/template-all")
def template_all(
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    if not branch_scope.sees_all_branches(current):
        raise HTTPException(403, {"code": "forbidden", "message": "ليس لديك صلاحية على كل الفروع."})
    try:
        made = sheet.apply_template_all(db, actor_user_id=current.id)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return {"created": made}


@router.post("/groups/copy")
def copy_groups(
    body: CopyIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    target = _branch(current, body.to_branch_id)
    source = _branch(current, body.from_branch_id)
    try:
        made = sheet.copy_groups(db, from_branch_id=source, to_branch_id=target,
                                 actor_user_id=current.id)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return {"created": [g.name for g in made]}


@router.get("/members")
def members(
    branch_id: int | None = Query(None),
    current: CurrentUser = Depends(require_capability(CAP_SALARY_VIEW)),
    db: Session = Depends(get_db),
) -> list[dict]:
    return sheet.members_of_branch(db, branch_id=_branch(current, branch_id))


@router.post("/members")
def assign(
    body: AssignIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    try:
        n = sheet.assign(db, branch_id=_branch(current, body.branch_id),
                         employee_ids=body.employee_ids, group_id=body.group_id,
                         actor_user_id=current.id)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return {"changed": n}


@router.post("/groups/{group_id}/members/order")
def order_members(
    group_id: int,
    body: OrderIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    _seen_group(db, group_id, current)
    sheet.reorder_members(db, group_id=group_id, employee_ids=body.ids)
    db.commit()
    return {"ok": True}


@router.get("/sheets")
def list_sheets(
    branch_id: int | None = Query(None),
    year: int | None = Query(None),
    current: CurrentUser = Depends(require_capability(CAP_SALARY_VIEW)),
    db: Session = Depends(get_db),
) -> list[dict]:
    if branch_scope.sees_all_branches(current):
        branch = branch_id if branch_id is not None else branch_scope.visible_branch_id(current)
    else:
        branch = _branch(current, branch_id)
    return sheet.list_sheets(db, branch_id=branch, year=year)


@router.get("/sheet")
def get_month(
    year: int,
    month: int,
    branch_id: int | None = Query(None),
    current: CurrentUser = Depends(require_capability(CAP_SALARY_VIEW)),
    db: Session = Depends(get_db),
) -> dict:
    branch = _branch(current, branch_id)
    run = sheet.current_run(db, branch_id=branch, year=year, month=month)
    if run is None or not sheet.is_sheet(db, run.id):
        return {"run": None, "groups": [], "summary": [], "totals": {}, "unassigned": [],
                "branch_id": branch}
    return sheet.sheet_out(db, run)


@router.get("/sheet/{run_id}")
def get_sheet(
    run_id: int,
    current: CurrentUser = Depends(require_capability(CAP_SALARY_VIEW)),
    db: Session = Depends(get_db),
) -> dict:
    return sheet.sheet_out(db, _seen_run(db, run_id, current))


@router.post("/sheet/prepare")
def prepare(
    body: MonthIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    try:
        run = sheet.prepare(db, branch_id=_branch(current, body.branch_id), year=body.year,
                            month=body.month, actor_user_id=current.id)
        out = sheet.sheet_out(db, run)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return out


def _row_out(db: Session, run: PayrollRun, employee_id: int) -> dict:
    data = sheet.sheet_out(db, run)
    for g in data["groups"]:
        for r in g["rows"]:
            if r["employee_id"] == employee_id:
                return {"group_id": g["id"], "row": r, "group_totals": g["totals"],
                        "totals": data["totals"], "summary": data["summary"],
                        "run": data["run"]}
    return {"run": data["run"]}


@router.put("/sheet/{run_id}/cell")
def set_cell(
    run_id: int,
    body: CellIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    run = _seen_run(db, run_id, current)
    try:
        sheet.set_cell(db, run_id=run_id, employee_id=body.employee_id, col_key=body.col_key,
                       value=body.value, actor_user_id=current.id)
        out = _row_out(db, run, body.employee_id)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return out


@router.put("/sheet/{run_id}/absence")
def set_absence(
    run_id: int,
    body: DaysIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    run = _seen_run(db, run_id, current)
    try:
        sheet.set_absent_days(db, run_id=run_id, employee_id=body.employee_id,
                              days=body.days, actor_user_id=current.id)
        out = _row_out(db, run, body.employee_id)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return out


@router.put("/sheet/{run_id}/note")
def set_note(
    run_id: int,
    body: NoteIn,
    current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
    db: Session = Depends(get_db),
) -> dict:
    _seen_run(db, run_id, current)
    try:
        sheet.set_note(db, run_id=run_id, employee_id=body.employee_id, notes=body.notes)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return {"ok": True}


def _act(db: Session, run_id: int, current: CurrentUser, fn) -> dict:
    run = _seen_run(db, run_id, current)
    try:
        result = fn(db, run_id=run_id, actor_user_id=current.id)
        db.flush()
        out = sheet.sheet_out(db, run)
        if isinstance(result, dict):
            out["result"] = result
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return out


class PayIn(BaseModel):
    treasury_id: int | None = None
    pay_date: date | None = None


@router.post("/sheet/{run_id}/pay")
def pay(run_id: int, body: PayIn,
        current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
        db: Session = Depends(get_db)) -> dict:
    def fn(db, *, run_id, actor_user_id):
        return sheet.pay(db, run_id=run_id, actor_user_id=actor_user_id,
                         treasury_id=body.treasury_id, pay_date=body.pay_date)
    return _act(db, run_id, current, fn)


@router.get("/sheet/{run_id}/payslip/{employee_id}")
def payslip(run_id: int, employee_id: int,
            current: CurrentUser = Depends(require_capability(CAP_SALARY_VIEW)),
            db: Session = Depends(get_db)) -> dict:
    _seen_run(db, run_id, current)
    try:
        return sheet.payslip(db, run_id=run_id, employee_id=employee_id)
    except _ERRORS as exc:
        _raise(exc)


@router.post("/sheet/{run_id}/close")
def close(run_id: int, current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
          db: Session = Depends(get_db)) -> dict:
    return _act(db, run_id, current, sheet.close)


@router.post("/sheet/{run_id}/reopen")
def reopen(run_id: int, current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
           db: Session = Depends(get_db)) -> dict:
    return _act(db, run_id, current, sheet.reopen)


@router.post("/sheet/{run_id}/post")
def post(run_id: int, current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
         db: Session = Depends(get_db)) -> dict:
    return _act(db, run_id, current, sheet.post)


@router.post("/sheet/{run_id}/reverse")
def reverse(run_id: int, current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
            db: Session = Depends(get_db)) -> dict:
    return _act(db, run_id, current, sheet.reverse)


@router.delete("/sheet/{run_id}")
def delete_sheet(run_id: int,
                 current: CurrentUser = Depends(require_capability(CAP_PAYROLL_POST)),
                 db: Session = Depends(get_db)) -> dict:
    _seen_run(db, run_id, current)
    try:
        sheet.delete_draft(db, run_id=run_id, actor_user_id=current.id)
    except _ERRORS as exc:
        _raise(exc)
    db.commit()
    return {"deleted": run_id}


@router.get("/sheet/{run_id}/export")
def export_xlsx(
    run_id: int,
    current: CurrentUser = Depends(require_capability(CAP_SALARY_VIEW)),
    db: Session = Depends(get_db),
):
    from fastapi.responses import StreamingResponse

    from src.models.org import Branch

    run = _seen_run(db, run_id, current)
    data = sheet.sheet_out(db, run)
    branch = db.get(Branch, run.branch_id) if run.branch_id else None
    buf = io.BytesIO()
    build_workbook(data, branch.name if branch else "").save(buf)
    buf.seek(0)
    name = f"payroll-{run.year}-{run.month:02d}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={name}"},
    )


MONTHS = ["يناير", "فبراير", "مارس", "ابريل", "مايو", "يونيو", "يوليو", "اغسطس", "سبتمبر",
          "اكتوبر", "نوفمبر", "ديسمبر"]


def build_workbook(data: dict, branch_name: str):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    run = data["run"]
    ws.title = f"{branch_name or 'المرتبات'}"[:31]
    ws.sheet_view.rightToLeft = True
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8F0E8")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    r = 1
    ws.cell(r, 1, f"مرتبات شهر {MONTHS[run['month'] - 1]} {run['year']}"
            + (f" — {branch_name}" if branch_name else "")).font = Font(bold=True, size=14)
    r += 2
    net_refs: list[tuple[str, str]] = []
    widest = 2
    for g in data["groups"]:
        cols = g["columns"]
        widest = max(widest, len(cols) + 4)
        earn = [c for c in cols if c["kind"] == "earning"]
        ded = [c for c in cols if c["kind"] != "earning"]
        order = earn + [{"key": "__earn", "label": "الاجمالي"}] + ded + [
            {"key": "__net", "label": "الصافي"}]
        ws.cell(r, 2, g["name"]).font = Font(bold=True, size=12)
        if earn:
            ws.cell(r, 3, "الاجور").font = bold
        if ded:
            ws.cell(r, 3 + len(earn) + 1, "الاستقطاعات").font = bold
        r += 1
        ws.cell(r, 1, "م")
        ws.cell(r, 2, "الاسم")
        for i, c in enumerate(order):
            ws.cell(r, 3 + i, c["label"])
        for i in range(1, len(order) + 3):
            cell = ws.cell(r, i)
            cell.font, cell.fill, cell.border, cell.alignment = bold, head_fill, box, center
        r += 1
        first = r
        col_of = {c["key"]: get_column_letter(3 + i) for i, c in enumerate(order)}
        for n, row in enumerate(g["rows"], start=1):
            ws.cell(r, 1, n)
            ws.cell(r, 2, row["name"])
            for c in earn + ded:
                ws.cell(r, 3 + order.index(c), float(Decimal(row["cells"][c["key"]]["value"])))
            if earn:
                ws.cell(r, 3 + order.index(order[len(earn)]),
                        "=" + "+".join(f"{col_of[c['key']]}{r}" for c in earn))
            else:
                ws.cell(r, 3 + len(earn), 0)
            minus = "".join(f"-{col_of[c['key']]}{r}" for c in ded)
            ws.cell(r, 3 + len(order) - 1, f"={col_of['__earn']}{r}{minus}")
            for i in range(1, len(order) + 3):
                ws.cell(r, i).border = box
                if i > 2:
                    ws.cell(r, i).number_format = "#,##0.00"
            r += 1
        last = r - 1
        ws.cell(r, 2, "الاجمالى").font = bold
        for i in range(len(order)):
            letter = get_column_letter(3 + i)
            ws.cell(r, 3 + i, f"=SUM({letter}{first}:{letter}{last})" if last >= first else 0)
        for i in range(1, len(order) + 3):
            cell = ws.cell(r, i)
            cell.font, cell.fill, cell.border = bold, head_fill, box
            if i > 2:
                cell.number_format = "#,##0.00"
        net_refs.append((g["name"], f"{col_of['__net']}{r}"))
        r += 3

    ws.cell(r, 2, "الملخص").font = Font(bold=True, size=12)
    r += 1
    start = r
    for name, ref in net_refs:
        ws.cell(r, 2, f"م {name}")
        ws.cell(r, 3, f"={ref}").number_format = "#,##0.00"
        ws.cell(r, 2).border = ws.cell(r, 3).border = box
        r += 1
    ws.cell(r, 2, "الاجمالى").font = bold
    ws.cell(r, 3, f"=SUM(C{start}:C{r - 1})" if r > start else 0).number_format = "#,##0.00"
    ws.cell(r, 3).font = bold
    ws.cell(r, 2).border = ws.cell(r, 3).border = box

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 24
    for i in range(3, widest + 3):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return wb
