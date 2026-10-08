from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=13)
_FILL = PatternFill("solid", fgColor="E6E6E6")
_THIN = Side(style="thin", color="888888")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_SOURCE = {"system": "النظام", "manual": "يدوي", "mixed": "النظام + يدوي", "computed": "محسوب"}


class _Sheet:
    def __init__(self, wb: Workbook, title: str, widths: list[int]):
        self.ws = wb.create_sheet(title[:31])
        self.ws.sheet_view.rightToLeft = True
        for i, w in enumerate(widths, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.r = 1

    def title(self, text: str):
        c = self.ws.cell(row=self.r, column=1, value=text)
        c.font = _TITLE
        self.r += 2

    def header(self, cols: list[str]):
        for i, t in enumerate(cols, start=1):
            c = self.ws.cell(row=self.r, column=i, value=t)
            c.font, c.fill, c.border = _BOLD, _FILL, _BOX
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        self.r += 1

    def row(self, vals: list, *, bold: bool = False, money_from: int = 2):
        for i, v in enumerate(vals, start=1):
            val = v
            if i >= money_from and v not in (None, "") and _is_num(str(v)):
                val = float(Decimal(str(v)))
            c = self.ws.cell(row=self.r, column=i, value=val)
            c.border = _BOX
            if isinstance(val, float):
                c.number_format = "#,##0.00"
            if bold:
                c.font, c.fill = _BOLD, _FILL
        self.r += 1

    def gap(self, n: int = 1):
        self.r += n


def _is_num(v: str) -> bool:
    try:
        Decimal(v)
        return True
    except Exception:
        return False


def _manual_rows(sh: _Sheet, rows: list[dict]):
    for ln in rows:
        sign = "−" if (ln.get("sign") or 1) < 0 else ""
        sh.row([ln["label"], ln.get("quantity") or "", ln.get("rate") or "",
                ln["signed"], f"يدوي {sign}".strip()])


def build(pkg: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    head = f"{pkg.get('branch_name') or ''} — {pkg['as_of']}"

    inv = pkg.get("inventory")
    if inv:
        for line in inv["lines"] + ([inv["unassigned"]] if inv.get("unassigned") else []):
            sh = _Sheet(wb, line["name"], [38, 14, 14, 14, 14, 16, 16, 16, 14])
            basis = "أصل السعر (سعر البيع)" if line["basis"] == "list" else "متوسط التكلفة"
            sh.title(f"{line['name']} — {head} — التقييم: {basis} × {line['factor_pct']}٪")
            sh.header(["الصنف", "الكود", "المخزن الرئيسي", "السيارات والمخازن", "سعر الوحدة",
                       "قيمة الرئيسي", "قيمة السيارات والمخازن", "القيمة", "مصدر السعر"])
            for it in line.get("items") or []:
                sh.row([it["name"], it["code"] or "", it["qty_main"], it["qty_other"],
                        it["unit_price"], it["value_main"], it["value_other"], it["value"],
                        it["price_source"]], money_from=3)
            sh.row(["الإجمالي", "", "", "", "", line["gross_main"], line["gross_other"],
                    line["gross"], ""], bold=True, money_from=3)
            sh.row([f"الصافي ({line['factor_pct']}٪)", "", "", "", "", line["net_main"],
                    line["net_other"], line["net_system"], ""], bold=True, money_from=3)
            if line["source"] == "manual":
                sh.row(["القيمة اليدوية المعتمدة", "", "", "", "", "", "", line["value"], "يدوي"],
                       bold=True, money_from=3)

        sh = _Sheet(wb, "قيمة المخازن", [34, 18, 18, 18, 14])
        sh.title(f"قيمة المخازن — {head}")
        sh.header(["خط الإنتاج", "المخزن الرئيسي", "السيارات والمخازن", "القيمة", "المصدر"])
        for line in inv["lines"]:
            sh.row([line["name"], line["net_main"], line["net_other"], line["value"],
                    _SOURCE.get(line["source"], line["source"])])
        for ln in inv.get("extras") or []:
            sh.row([ln["label"], "", "", ln["signed"], "يدوي"])
        sh.row(["قيمة المخازن", "", "", inv["total"], ""], bold=True)
        if inv.get("unassigned"):
            sh.gap()
            count = inv["unassigned"]["items_count"]
            sh.row([f"أصناف خارج الخطوط (غير مدرجة في الإجمالي): {count}", "",
                    "", inv["unassigned"]["gross"], ""])

    bal = pkg.get("balances")
    if bal:
        sh = _Sheet(wb, "عملاء ونقدية وموردين", [38, 14, 14, 18, 14])
        sh.title(f"عملاء ونقدية وموردين — {head}")

        def block(title, rows, total, *, label_key="name"):
            sh.row([title], bold=True)
            sh.header(["البند", "", "", "المبلغ", "المصدر"])
            for r in rows:
                sh.row([r.get(label_key) or r.get("label"), "", "", r["amount"], "النظام"])

        c = bal["customers"]
        block("مديونية العملاء", c["rows"], c["total"], label_key="label")
        _manual_rows(sh, c["adjustments"])
        sh.row(["الإجمالي", "", "", c["total"], ""], bold=True)
        sh.gap()
        for key, title in (("suppliers", "مستحقات الموردين"), ("accrued", "مصروفات مستحقة"),
                           ("partners", "جاري الشركاء"), ("cash", "النقدية")):
            b = bal[key]
            block(title, b["rows"], b["total"])
            _manual_rows(sh, b["manual"])
            sh.row(["الإجمالي", "", "", b["total"], ""], bold=True)
            sh.gap()
        for key, title in (("technicians", "مستحقات الفنيين"), ("gifts", "رصيد الهدايا")):
            b = bal[key]
            sh.row([title], bold=True)
            sh.header(["البند", "الكمية", "السعر", "المبلغ", "المصدر"])
            _manual_rows(sh, b["rows"])
            sh.row(["الإجمالي" if key == "technicians" else "الصافي", "", "", b["total"], ""],
                   bold=True)
            sh.gap()
        bo = bal["bonuses"]
        sh.row(["بوانص التجار"], bold=True)
        sh.header(["السيارة", *bo["years"], "الإجمالي"])
        for r in bo["rows"]:
            sh.row([r["label"], *[r["years"].get(y, "") for y in bo["years"]], r["total"]])
        sh.row(["الإجمالي", *[bo["by_year"].get(y, "") for y in bo["years"]], bo["total"]],
               bold=True)
        sh.gap()
        sh.row(["تكلفة البوانص"], bold=True)
        for cl in bo["calc"]:
            sh.row([cl["label"], "", "", cl["value"], ""])
        sh.row(["الصافي (قيمة بوانص التجار)", "", "", bo["net"], ""], bold=True)

    bs = pkg.get("balance_sheet")
    if bs:
        sh = _Sheet(wb, "الميزانية العمومية", [36, 18, 18, 14])
        sh.title(f"ميزانية عمومية في {pkg['as_of']} — {pkg.get('branch_name') or ''}")
        sh.header(["الأصول", "فرعي", "رئيسي", "المصدر"])
        a = bs["assets"]
        sh.row(["الأصول المتداولة", "", a["current_total"], ""], bold=True)
        for x in a["current"]:
            sh.row([x["label"], x["amount"], "", _SOURCE.get(x["source"], "")])
        sh.row(["الأصول الثابتة", "", a["fixed_total"], ""], bold=True)
        for x in a["fixed"]:
            sh.row([x["label"], x["amount"], "", _SOURCE.get(x["source"], "")])
        sh.row(["أرصدة مدينة أخرى", "", a["other_total"], ""], bold=True)
        for x in a["other"]:
            sh.row([x["label"], x["amount"], "", _SOURCE.get(x["source"], "")])
        sh.row(["إجمالي الأصول", "", a["total"], ""], bold=True)
        sh.gap()
        sh.header(["الخصوم وحقوق الملكية", "فرعي", "رئيسي", "المصدر"])
        sh.row(["حقوق الملكية", "", bs["equity"]["total"], ""], bold=True)
        for x in bs["equity"]["rows"]:
            sh.row([x["label"], x["amount"], "", _SOURCE.get(x["source"], "")])
        sh.row(["الالتزامات", "", bs["liabilities"]["total"], ""], bold=True)
        for x in bs["liabilities"]["rows"]:
            sh.row([x["label"], x["amount"], "", _SOURCE.get(x["source"], "")])
        sh.row(["الإجمالي", "", bs["total"], ""], bold=True)
        for x in bs.get("memo") or []:
            sh.row([f"{x['label']} (للعلم)", x["amount"], "", "يدوي"])

    st = pkg.get("settlement")
    if st and st["areas"]:
        sh = _Sheet(wb, "تسوية المناطق", [40, 18, 14])
        sh.title(f"تسوية المناطق — {head}")
        for area in st["areas"]:
            sh.row([area["name"]], bold=True)
            sh.row(["المديونية", area["debt"], "النظام"])
            for d in area["deductions"]:
                sh.row([d["label"], d["signed"], "يدوي"])
            sh.row(["الصافي المستحق على المنطقة", area["net"], ""], bold=True)
            sh.gap()

    adv = pkg.get("advances")
    if adv:
        dates = adv["dates"]
        sh = _Sheet(wb, "سلف الموظفين", [36] + [16] * len(dates))
        sh.title(f"سلف الموظفين — {pkg.get('branch_name') or ''}")
        sh.header(["الموظف / الحساب", *dates])
        for r in adv["rows"]:
            sh.row([r["label"], *[r["values"].get(d, "") for d in dates]])
        sh.row(["الإجمالي", *[adv["totals"].get(d) for d in dates]], bold=True)

    if not wb.sheetnames:
        wb.create_sheet("فارغ")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
